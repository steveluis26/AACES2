"""Importar participantes de un grupo desde Excel (.xlsx) o CSV.

Flujo en dos pasos con el MISMO archivo: vista previa (no guarda nada) y
confirmación. Reconoce encabezados con nombres distintos ("Correo electrónico",
"Celular", "Puesto"…), valida cada fila y detecta duplicados dentro del archivo
y contra la organización (misma CURP o mismo correo = misma persona, se reutiliza).
Inscribir no da folio ni QR ni gasta constancias (eso ocurre al acreditar).
"""
import csv
import io
import re
import unicodedata
import uuid
from typing import Any, Dict, List, Optional, Tuple

from fastapi import HTTPException
from sqlalchemy import text
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.cifrado import cifrar, indice

MAX_BYTES = 5 * 1024 * 1024
MAX_FILAS = 1000

CAMPOS = ["nombres", "apellido_paterno", "apellido_materno", "correo", "telefono", "curp", "ocupacion", "puesto", "empresa", "ciudad"]
ETIQUETAS = {
    "nombres": "Nombre(s)", "apellido_paterno": "Apellido paterno", "apellido_materno": "Apellido materno",
    "correo": "Correo", "telefono": "Teléfono", "curp": "CURP", "ocupacion": "Ocupación específica",
    "puesto": "Puesto", "empresa": "Empresa", "ciudad": "Ciudad",
}

# Encabezados que reconocemos (ya normalizados: minúsculas, sin acentos ni signos)
SINONIMOS = {
    "nombres": {"nombre", "nombres", "nombre s", "nombres del trabajador", "nombre del trabajador", "primer nombre"},
    "nombre_completo": {"nombre completo", "nombre y apellidos", "nombre completo del trabajador", "trabajador", "participante", "alumno", "nombre del participante"},
    "apellido_paterno": {"apellido paterno", "paterno", "primer apellido", "ap paterno"},
    "apellido_materno": {"apellido materno", "materno", "segundo apellido", "ap materno"},
    "apellidos": {"apellidos", "apellido"},
    "correo": {"correo", "email", "e mail", "correo electronico", "mail", "correo e", "e-mail"},
    "telefono": {"telefono", "celular", "tel", "movil", "whatsapp", "numero de telefono", "telefono celular"},
    "curp": {"curp", "clave unica de registro de poblacion"},
    "ocupacion": {"ocupacion", "ocupacion especifica", "ocupacion especifica catalogo nacional de ocupaciones"},
    "puesto": {"puesto", "cargo", "puesto de trabajo", "area"},
    "empresa": {"empresa", "razon social", "compania", "empresa donde labora", "patron"},
    "ciudad": {"ciudad", "municipio", "localidad", "ciudad de origen"},
}

CORREO = re.compile(r"^[^@\s]+@[^@\s]+\.[^@\s]+$")


def _norm(v: Any) -> str:
    s = unicodedata.normalize("NFD", str(v or "")).encode("ascii", "ignore").decode().lower()
    return re.sub(r"\s+", " ", re.sub(r"[^a-z0-9]+", " ", s)).strip()


def _celda(v: Any) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)  # teléfonos que Excel guarda como número
    return re.sub(r"\s+", " ", str(v)).strip()


def leer_archivo(contenido: bytes, nombre: str) -> List[List[str]]:
    if len(contenido) > MAX_BYTES:
        raise HTTPException(400, "El archivo pesa más de 5 MB")
    ext = (nombre or "").lower().rsplit(".", 1)[-1]
    if ext == "xlsx":
        try:
            from openpyxl import load_workbook
            wb = load_workbook(io.BytesIO(contenido), read_only=True, data_only=True)
        except Exception:
            raise HTTPException(400, "No pudimos abrir el Excel. Guárdalo como .xlsx e inténtalo de nuevo.")
        hoja = wb.worksheets[0]
        filas = [[_celda(c) for c in fila] for fila in hoja.iter_rows(values_only=True)]
        wb.close()
    elif ext == "csv":
        texto = None
        for enc in ("utf-8-sig", "latin-1"):
            try:
                texto = contenido.decode(enc)
                break
            except UnicodeDecodeError:
                continue
        muestra = (texto or "")[:2000]
        delim = ";" if muestra.count(";") > muestra.count(",") else ","
        filas = [[_celda(c) for c in fila] for fila in csv.reader(io.StringIO(texto or ""), delimiter=delim)]
    elif ext == "xls":
        raise HTTPException(400, "Ese es el formato viejo de Excel (.xls). Ábrelo y guárdalo como .xlsx.")
    else:
        raise HTTPException(400, "Sube un archivo de Excel (.xlsx) o CSV")
    return [f for f in filas if any(c for c in f)]


def _mapear(encabezados: List[str]) -> Tuple[Dict[int, str], List[str]]:
    """Índice de columna -> campo. Devuelve también los encabezados que no se usaron."""
    mapa, ignorados, usados = {}, [], set()
    for i, h in enumerate(encabezados):
        n = _norm(h)
        campo = next((c for c, sin in SINONIMOS.items() if n in sin), None)
        if campo and campo not in usados:
            mapa[i] = campo
            usados.add(campo)
        elif h:
            ignorados.append(h)
    return mapa, ignorados


def _fila_encabezados(filas: List[List[str]]) -> int:
    """La primera fila (de las 10 primeras) con al menos dos encabezados reconocidos."""
    for i, f in enumerate(filas[:10]):
        if len(_mapear(f)[0]) >= 2:
            return i
    raise HTTPException(400, "No encontramos los encabezados. Usa la plantilla o pon en la primera fila columnas como Nombre, Apellido paterno y CURP.")


def _datos_fila(f: List[str], mapa: Dict[int, str]) -> Dict[str, str]:
    d = {c: "" for c in CAMPOS}
    extra = {}
    for i, campo in mapa.items():
        v = f[i] if i < len(f) else ""
        if campo in ("nombre_completo", "apellidos"):
            extra[campo] = v
        else:
            d[campo] = v
    if not d["nombres"] and extra.get("nombre_completo"):
        d["nombres"] = extra["nombre_completo"]  # se guarda completo en "nombre"
    if extra.get("apellidos") and not d["apellido_paterno"]:
        partes = extra["apellidos"].split(" ", 1)
        d["apellido_paterno"] = partes[0]
        d["apellido_materno"] = partes[1] if len(partes) > 1 else ""
    d["curp"] = re.sub(r"[\s-]", "", d["curp"]).upper()
    d["correo"] = d["correo"].lower()
    return d


async def analizar(db: AsyncSession, curso_id: str, cliente_id: str, contenido: bytes, nombre: str) -> Dict[str, Any]:
    filas = leer_archivo(contenido, nombre)
    if not filas:
        raise HTTPException(400, "El archivo está vacío")
    i_enc = _fila_encabezados(filas)
    mapa, ignorados = _mapear(filas[i_enc])
    datos = filas[i_enc + 1:]
    if not datos:
        raise HTTPException(400, "El archivo solo tiene encabezados")
    if len(datos) > MAX_FILAS:
        raise HTTPException(400, f"Máximo {MAX_FILAS} participantes por archivo")

    # Lo que ya existe en la organización (para reutilizar) y en el grupo (para omitir)
    existentes = (await db.execute(text("""
        SELECT p.id, p.curp_hash, lower(p.correo),
               EXISTS (SELECT 1 FROM aaces.curso_participante cp WHERE cp.participante_id = p.id AND cp.curso_id = :curso)
        FROM aaces.participantes p WHERE p.cliente_id = :cid
    """), {"cid": cliente_id, "curso": curso_id})).fetchall()
    por_curp = {r[1]: (str(r[0]), r[3]) for r in existentes if r[1]}
    por_correo = {r[2]: (str(r[0]), r[3]) for r in existentes if r[2]}
    # Sin CURP ni correo solo queda el nombre: se compara contra los inscritos del grupo
    # (no contra toda la agencia, para no confundir homónimos) y así reimportar no duplica
    inscritos_nombre = {
        _norm(" ".join(x for x in r if x)) for r in (await db.execute(text("""
            SELECT p.nombre, COALESCE(p.apellido_paterno, p.apellido), p.apellido_materno
            FROM aaces.curso_participante cp JOIN aaces.participantes p ON p.id = cp.participante_id
            WHERE cp.curso_id = :curso
        """), {"curso": curso_id})).fetchall()
    }

    resultado, vistos_curp, vistos_correo = [], {}, {}
    for n, f in enumerate(datos, start=i_enc + 2):  # número de fila como lo ve Excel
        d = _datos_fila(f, mapa)
        errores, avisos = [], []
        if not d["nombres"]:
            errores.append("Falta el nombre")
        if d["curp"] and (len(d["curp"]) != 18 or not d["curp"].isalnum()):
            errores.append("La CURP debe tener 18 caracteres")
        if d["correo"] and not CORREO.match(d["correo"]):
            errores.append("Correo no válido")
        if d["curp"] and d["curp"] in vistos_curp:
            errores.append(f"CURP repetida en la fila {vistos_curp[d['curp']]}")
        if d["correo"] and d["correo"] in vistos_correo:
            errores.append(f"Correo repetido en la fila {vistos_correo[d['correo']]}")
        if d["curp"]:
            vistos_curp.setdefault(d["curp"], n)
        if d["correo"]:
            vistos_correo.setdefault(d["correo"], n)

        previo = (por_curp.get(indice(d["curp"])) if d["curp"] else None) or (por_correo.get(d["correo"]) if d["correo"] else None)
        nombre_completo = _norm(" ".join(x for x in (d["nombres"], d["apellido_paterno"], d["apellido_materno"]) if x))
        if errores:
            estado = "error"
        elif not previo and not d["curp"] and not d["correo"] and nombre_completo in inscritos_nombre:
            estado = "inscrito"
            avisos.append("Parece que ya está inscrito (mismo nombre). Agrega su CURP para distinguirlo")
        elif previo and previo[1]:
            estado = "inscrito"
            avisos.append("Ya está inscrito en este grupo")
        elif previo:
            estado = "existente"
            avisos.append("Ya lo tienes registrado: se reutiliza su expediente")
        else:
            estado = "nuevo"
        if not errores and not d["curp"]:
            avisos.append("Sin CURP (la necesitarás para el DC-3)")
        resultado.append({"fila": n, "datos": d, "estado": estado, "errores": errores, "avisos": avisos,
                          "participante_id": previo[0] if previo else None})

    resumen = {e: sum(1 for r in resultado if r["estado"] == e) for e in ("nuevo", "existente", "inscrito", "error")}
    return {
        "columnas": [{"encabezado": filas[i_enc][i], "campo": c, "etiqueta": ETIQUETAS.get(c, "Nombre completo" if c == "nombre_completo" else "Apellidos")} for i, c in mapa.items()],
        "ignoradas": ignorados,
        "filas": resultado,
        "resumen": {**resumen, "total": len(resultado), "a_importar": resumen["nuevo"] + resumen["existente"]},
    }


async def importar(db: AsyncSession, curso_id: str, cliente_id: str, analisis: Dict[str, Any]) -> Dict[str, int]:
    """Crea o reutiliza a cada participante válido y lo inscribe en el grupo (sin folio)."""
    creados = reutilizados = 0
    for r in analisis["filas"]:
        if r["estado"] not in ("nuevo", "existente"):
            continue
        d = r["datos"]
        pid = r["participante_id"]
        if pid:
            # Solo completa lo que falte; no sobreescribe lo que la agencia ya capturó
            await db.execute(text("""
                UPDATE aaces.participantes SET
                  correo = COALESCE(NULLIF(correo, ''), :correo),
                  telefono = COALESCE(NULLIF(telefono, ''), :telefono),
                  empresa = COALESCE(NULLIF(empresa, ''), :empresa),
                  cargo = COALESCE(NULLIF(cargo, ''), :puesto),
                  ciudad_origen = COALESCE(NULLIF(ciudad_origen, ''), :ciudad),
                  ocupacion = COALESCE(NULLIF(ocupacion, ''), :ocupacion),
                  curp = COALESCE(curp, :curp), curp_hash = COALESCE(curp_hash, :curp_hash)
                WHERE id = CAST(:id AS uuid)
            """), {"id": pid, "correo": d["correo"] or None, "telefono": d["telefono"] or None, "empresa": d["empresa"] or None,
                   "puesto": d["puesto"] or None, "ciudad": d["ciudad"] or None, "ocupacion": d["ocupacion"][:150] or None,
                   "curp": cifrar(d["curp"]), "curp_hash": indice(d["curp"])})
            reutilizados += 1
        else:
            apellidos = " ".join(x for x in (d["apellido_paterno"], d["apellido_materno"]) if x)
            pid = str((await db.execute(text("""
                INSERT INTO aaces.participantes (id, pax_id, cliente_id, nombre, apellido, apellido_paterno, apellido_materno,
                    correo, telefono, empresa, cargo, ciudad_origen, ocupacion, curp, curp_hash, pais)
                VALUES (gen_random_uuid(), :pax, :cid, :nombre, :apellido, :ap_pat, :ap_mat,
                    :correo, :telefono, :empresa, :puesto, :ciudad, :ocupacion, :curp, :curp_hash, 'Mexico')
                RETURNING id
            """), {"pax": f"PAX-{uuid.uuid4().hex[:8].upper()}", "cid": cliente_id, "nombre": d["nombres"][:100],
                   "apellido": apellidos or None, "ap_pat": d["apellido_paterno"] or None, "ap_mat": d["apellido_materno"] or None,
                   "correo": d["correo"] or None, "telefono": d["telefono"] or None, "empresa": d["empresa"] or None,
                   "puesto": d["puesto"] or None, "ciudad": d["ciudad"] or None, "ocupacion": d["ocupacion"][:150] or None,
                   "curp": cifrar(d["curp"]), "curp_hash": indice(d["curp"])})).scalar())
            creados += 1
        await db.execute(text("""
            INSERT INTO aaces.curso_participante (id, curso_id, participante_id, estado_pago, estado_acreditacion,
                fecha_inicio_vigencia, valor_pagado, costo_asignado, descuento)
            SELECT gen_random_uuid(), c.id, CAST(:pid AS uuid), 'pendiente', false,
                   COALESCE(c.fecha_fin, c.fecha_inicio, CURRENT_DATE), 0,
                   COALESCE(c.precio_base, c.precio_promocional, c.costo_total, 0), 0
            FROM aaces.cursos c
            WHERE c.id = CAST(:curso AS uuid)
              AND NOT EXISTS (SELECT 1 FROM aaces.curso_participante x WHERE x.curso_id = c.id AND x.participante_id = CAST(:pid AS uuid))
        """), {"pid": pid, "curso": curso_id})
    await db.commit()
    return {"creados": creados, "reutilizados": reutilizados, "inscritos": creados + reutilizados}


def plantilla_xlsx() -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    wb = Workbook()
    ws = wb.active
    ws.title = "Participantes"
    encabezados = [ETIQUETAS[c] for c in CAMPOS]
    ws.append(encabezados)
    ws.append(["María Fernanda", "López", "Ruiz", "maria.lopez@empresa.com", "4421234567", "LORM900101MQTPZR05",
               "Supervisora de seguridad", "Supervisora", "Grupo Industrial del Bajío", "Querétaro"])
    naranja = PatternFill("solid", fgColor="F97316")
    for i, celda in enumerate(ws[1], start=1):
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = naranja
        celda.alignment = Alignment(vertical="center")
        ws.column_dimensions[celda.column_letter].width = max(16, len(encabezados[i - 1]) + 6)
    ws.column_dimensions["F"].width = 24
    for celda in ws[2]:
        celda.font = Font(italic=True, color="888888")
    ws.freeze_panes = "A2"
    ayuda = wb.create_sheet("Instrucciones")
    for linea in [
        "Cómo llenar la plantilla",
        "",
        "• Un trabajador por fila. Borra la fila de ejemplo (en gris).",
        "• Obligatorio: Nombre(s). Recomendado: Apellido paterno, Apellido materno y CURP (el DC-3 los pide).",
        "• La CURP debe tener 18 caracteres. El correo y el teléfono son opcionales.",
        "• Si un trabajador ya existe (misma CURP o correo) se reutiliza su expediente.",
        "• Inscribir no gasta constancias: los folios se generan al acreditar al grupo.",
    ]:
        ayuda.append([linea])
    ayuda["A1"].font = Font(bold=True, size=13)
    ayuda.column_dimensions["A"].width = 100
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
